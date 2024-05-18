import streamlit as st
from pptx import Presentation
from io import BytesIO
from docx import Document
from transformers import pipeline

# Load pretrained summarization model
summarizer = pipeline("summarization")

def extract_text_from_ppt(ppt):
    pr=Presentation(ppt)
    text=[]
    for slide in pr.slides:
        for shape in slide.shapes:
            if hasattr(shape,"text"):
                text.append(shape.text)
    return "\n".join(text)


def extract_text_from_docx(file):
    text = ""
    doc = Document(file)
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

# Streamlit app
st.set_page_config(page_title="Document text summarizer", page_icon=":bar_chart:",layout="wide")
st.title(":bar_chart: Document Text Summarizer")
st.image

uploaded_file = st.file_uploader("Upload a PPT or Word document", type=["pptx", "docx"])

if uploaded_file:
    file_ext = uploaded_file.name.split(".")[-1]
    if file_ext == "pptx":
        extracted_text = extract_text_from_ppt(uploaded_file)
    elif file_ext == "docx":
        extracted_text = extract_text_from_docx(uploaded_file)

    st.write("Extracted Text:")
    st.write(extracted_text)
    summary = summarizer(extracted_text, max_length=150, min_length=30, do_sample=False)

    st.write("Summary:")
    st.write(summary[0]['summary_text'])
